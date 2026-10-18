# SPDX-FileCopyrightText: 2026 Thomas Faust
# SPDX-License-Identifier: LicenseRef-Docuvate-SUL-1.0

from __future__ import annotations

import csv
import io
import zipfile


def synthetic_docx_bytes() -> bytes:
    from docx import Document  # noqa: PLC0415

    doc = Document()
    doc.add_heading("Synthetic Memo", level=1)
    doc.add_paragraph("Invoice total EUR 199,00 for regression testing.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def synthetic_xlsx_bytes() -> bytes:
    from openpyxl import Workbook  # noqa: PLC0415

    wb = Workbook()
    ws = wb.active
    ws.title = "Costs"
    ws.append(["Item", "Amount"])
    ws.append(["Widget", "199"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def synthetic_csv_bytes() -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Item", "Amount"])
    writer.writerow(["Widget", "199"])
    return buf.getvalue().encode("utf-8")


def synthetic_pptx_bytes() -> bytes:
    from pptx import Presentation  # noqa: PLC0415

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Synthetic Slide"
    slide.placeholders[1].text = "Pipeline smoke slide body."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def synthetic_md_bytes() -> bytes:
    return b"# Synthetic Markdown\n\nOffice ingest **bold** line.\n"


def minimal_odt_bytes() -> bytes:
    """Minimal ODT (zip + mimetype) for odfpy load smoke."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        zf.writestr(
            "content.xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
  <office:body><office:text><text:p>Synthetic ODT paragraph.</text:p></office:text></office:body>
</office:document-content>""",
        )
    return buf.getvalue()
