# SPDX-FileCopyrightText: 2026 Thomas Faust
# SPDX-License-Identifier: LicenseRef-Docuvate-SUL-1.0

from __future__ import annotations

import csv
import email
import io
import logging
import re
from html.parser import HTMLParser
from typing import TYPE_CHECKING

from docuvate_worker.domain.layout_ir import (
    LAYOUT_IR_VERSION,
    LayoutIrBlock,
    LayoutIrDocument,
    LayoutIrPage,
)
from docuvate_worker.domain.models import ExtractionBlock
from docuvate_worker.infrastructure.extractors.pdf_native import try_extract_pdf_native
from docuvate_worker.infrastructure.office.layout_tables import layout_ir_from_grid
from docuvate_worker.infrastructure.office.libreoffice import convert_bytes_to_pdf
from docuvate_worker.infrastructure.office.mimes import (
    LEGACY_OFFICE_MIMES,
    normalize_mime,
)
from docuvate_worker.infrastructure.office.types import OfficeExtractResult

if TYPE_CHECKING:
    from collections.abc import Sequence

logger = logging.getLogger(__name__)

_BLOCK_INDEX = 0


def _next_block_index() -> int:
    global _BLOCK_INDEX  # noqa: PLW0603
    _BLOCK_INDEX += 1
    return _BLOCK_INDEX


def _paragraph_blocks(paragraphs: Sequence[str], *, page: int = 1) -> tuple[ExtractionBlock, ...]:
    blocks: list[ExtractionBlock] = []
    y = 0.05
    for para in paragraphs:
        text = para.strip()
        if not text:
            y += 0.02
            continue
        height = min(0.12, 0.03 + text.count("\n") * 0.015)
        blocks.append(
            ExtractionBlock(
                page=page,
                x=0.08,
                y=y,
                width=0.84,
                height=height,
                text=text,
                block_index=_next_block_index(),
            )
        )
        y += height + 0.02
        if y > 0.95:
            break
    return tuple(blocks)


def _result_from_text(text: str, *, source: str) -> OfficeExtractResult:
    paragraphs = [p for p in re.split(r"\n\s*\n", text) if p.strip()]
    if not paragraphs:
        paragraphs = [line for line in text.splitlines() if line.strip()]
    blocks = _paragraph_blocks(paragraphs)
    joined = "\n\n".join(p.strip() for p in paragraphs)
    return OfficeExtractResult(text=joined.strip(), blocks=blocks, source=source)


def _extract_docx(content: bytes) -> OfficeExtractResult:
    from docx import Document  # noqa: PLC0415

    doc = Document(io.BytesIO(content))
    paragraphs: list[str] = []
    for para in doc.paragraphs:
        raw = para.text.strip()
        if not raw:
            continue
        style_name = para.style.name if para.style is not None else ""
        if style_name.startswith("Heading"):
            paragraphs.append(raw.upper())
        else:
            paragraphs.append(raw)
    return _result_from_text("\n\n".join(paragraphs), source="docx")


def _sheet_rows_from_openpyxl(content: bytes) -> list[tuple[str, list[list[str]]]]:
    from openpyxl import load_workbook  # noqa: PLC0415

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheets: list[tuple[str, list[list[str]]]] = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                rows.append(cells)
        if rows:
            sheets.append((name, rows))
    wb.close()
    return sheets


def _extract_xlsx(content: bytes) -> OfficeExtractResult:
    sheets = _sheet_rows_from_openpyxl(content)
    if not sheets:
        return OfficeExtractResult(text="", blocks=(), source="xlsx")
    text_parts: list[str] = []
    layout_pages: list[LayoutIrPage] = []
    all_blocks: list[ExtractionBlock] = []
    for page_idx, (name, rows) in enumerate(sheets, start=1):
        text_parts.append(f"## {name}\n")
        for row in rows:
            text_parts.append("\t".join(row))
        sheet_doc = layout_ir_from_grid(name, rows, page=page_idx)
        if sheet_doc.pages:
            layout_pages.append(sheet_doc.pages[0])
        y = 0.05
        for row in rows[:40]:
            line = "\t".join(row)
            all_blocks.append(
                ExtractionBlock(
                    page=page_idx,
                    x=0.05,
                    y=y,
                    width=0.9,
                    height=0.035,
                    text=line,
                    block_index=_next_block_index(),
                )
            )
            y += 0.04
    layout = LayoutIrDocument(version=LAYOUT_IR_VERSION, pages=tuple(layout_pages))
    return OfficeExtractResult(
        text="\n".join(text_parts).strip(),
        blocks=tuple(all_blocks),
        layout_ir=layout,
        source="xlsx",
    )


def _extract_pptx(content: bytes) -> OfficeExtractResult:
    from pptx import Presentation  # noqa: PLC0415

    prs = Presentation(io.BytesIO(content))
    blocks: list[ExtractionBlock] = []
    text_parts: list[str] = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_lines: list[str] = []
        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue
            text = str(shape.text).strip()
            if text:
                slide_lines.append(text)
        if not slide_lines:
            continue
        text_parts.append(f"## Slide {slide_idx}\n" + "\n".join(slide_lines))
        y = 0.1
        for line in slide_lines:
            blocks.append(
                ExtractionBlock(
                    page=slide_idx,
                    x=0.08,
                    y=y,
                    width=0.84,
                    height=0.08,
                    text=line,
                    block_index=_next_block_index(),
                )
            )
            y += 0.1
    return OfficeExtractResult(
        text="\n\n".join(text_parts).strip(),
        blocks=tuple(blocks),
        source="pptx",
    )


def _extract_odt(content: bytes) -> OfficeExtractResult:
    from odf import teletype  # noqa: PLC0415
    from odf.opendocument import load  # noqa: PLC0415

    doc = load(io.BytesIO(content))
    text = teletype.extractText(doc)
    return _result_from_text(text, source="odt")


def _extract_ods(content: bytes) -> OfficeExtractResult:
    from odf import teletype  # noqa: PLC0415
    from odf.opendocument import load  # noqa: PLC0415
    from odf.table import Table, TableCell, TableRow  # noqa: PLC0415
    from odf.text import P  # noqa: PLC0415

    doc = load(io.BytesIO(content))
    rows: list[list[str]] = []
    for table in doc.getElementsByType(Table):
        for row_el in table.getElementsByType(TableRow):
            row: list[str] = []
            for cell in row_el.getElementsByType(TableCell):
                parts = [teletype.extractText(p) for p in cell.getElementsByType(P)]
                row.append(" ".join(p for p in parts if p).strip())
            if any(row):
                rows.append(row)
    if not rows:
        return OfficeExtractResult(text="", blocks=(), source="ods")
    layout = layout_ir_from_grid("Sheet1", rows)
    text = "\n".join("\t".join(r) for r in rows)
    blocks = _paragraph_blocks(text.splitlines())
    return OfficeExtractResult(text=text, blocks=blocks, layout_ir=layout, source="ods")


def _extract_odp(content: bytes) -> OfficeExtractResult:
    from odf import teletype  # noqa: PLC0415
    from odf.draw import Page as DrawPage  # noqa: PLC0415
    from odf.opendocument import load  # noqa: PLC0415
    from odf.text import P  # noqa: PLC0415

    doc = load(io.BytesIO(content))
    blocks: list[ExtractionBlock] = []
    text_parts: list[str] = []
    for page_idx, page in enumerate(doc.getElementsByType(DrawPage), start=1):
        lines = [teletype.extractText(p) for p in page.getElementsByType(P)]
        lines = [ln.strip() for ln in lines if ln.strip()]
        if not lines:
            continue
        text_parts.append(f"## Slide {page_idx}\n" + "\n".join(lines))
        y = 0.1
        for line in lines:
            blocks.append(
                ExtractionBlock(
                    page=page_idx,
                    x=0.08,
                    y=y,
                    width=0.84,
                    height=0.08,
                    text=line,
                    block_index=_next_block_index(),
                )
            )
            y += 0.1
    return OfficeExtractResult(
        text="\n\n".join(text_parts).strip(),
        blocks=tuple(blocks),
        source="odp",
    )


def _extract_rtf(content: bytes) -> OfficeExtractResult:
    from striprtf.striprtf import rtf_to_text  # noqa: PLC0415

    text = rtf_to_text(content.decode("utf-8", errors="replace"))
    return _result_from_text(text, source="rtf")


def _extract_csv(content: bytes) -> OfficeExtractResult:
    decoded = content.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(decoded))
    rows = [[cell.strip() for cell in row] for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        return OfficeExtractResult(text="", blocks=(), source="csv")
    layout = layout_ir_from_grid("CSV", rows)
    text = "\n".join("\t".join(r) for r in rows)
    blocks = _paragraph_blocks(text.splitlines())
    return OfficeExtractResult(text=text, blocks=blocks, layout_ir=layout, source="csv")


def _extract_plain(content: bytes, *, source: str) -> OfficeExtractResult:
    text = content.decode("utf-8", errors="replace")
    return _result_from_text(text, source=source)


class _HtmlTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._chunks: list[str] = []
        self._skip = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self._skip = True
        if tag in {"p", "br", "div", "h1", "h2", "h3", "li", "tr"}:
            self._chunks.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"}:
            self._skip = False

    def handle_data(self, data: str) -> None:
        if self._skip:
            return
        chunk = data.strip()
        if chunk:
            self._chunks.append(chunk)

    def text(self) -> str:
        return re.sub(r"\n{3,}", "\n\n", "".join(self._chunks))


def _extract_html(content: bytes) -> OfficeExtractResult:
    parser = _HtmlTextParser()
    parser.feed(content.decode("utf-8", errors="replace"))
    return _result_from_text(parser.text(), source="html")


def _extract_eml(content: bytes) -> OfficeExtractResult:
    msg = email.message_from_bytes(content)
    parts: list[str] = []
    subject = msg.get("Subject")
    if subject:
        parts.append(f"Subject: {subject}")
    from_hdr = msg.get("From")
    if from_hdr:
        parts.append(f"From: {from_hdr}")
    body = _email_body_text(msg)
    if body:
        parts.append(body)
    return _result_from_text("\n\n".join(parts), source="eml")


def _email_body_text(msg: email.message.Message) -> str:
    if msg.is_multipart():
        texts: list[str] = []
        for part in msg.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True)
                if isinstance(payload, bytes):
                    texts.append(payload.decode("utf-8", errors="replace"))
        return "\n".join(texts).strip()
    payload = msg.get_payload(decode=True)
    if isinstance(payload, bytes):
        return payload.decode("utf-8", errors="replace")
    if isinstance(payload, str):
        return payload
    return ""


def _extract_msg(content: bytes) -> OfficeExtractResult:
    import extract_msg  # noqa: PLC0415

    msg = extract_msg.Message(io.BytesIO(content))
    parts: list[str] = []
    if msg.subject:
        parts.append(f"Subject: {msg.subject}")
    if msg.sender:
        parts.append(f"From: {msg.sender}")
    if msg.body:
        parts.append(str(msg.body))
    msg.close()
    return _result_from_text("\n\n".join(parts), source="msg")


def _extract_via_libreoffice_to_pdf(content: bytes, mime_type: str) -> OfficeExtractResult | None:
    pdf = convert_bytes_to_pdf(content, mime_type)
    if pdf is None:
        return None
    native = try_extract_pdf_native(pdf)
    if native is None:
        return None
    text, blocks = native
    return OfficeExtractResult(
        text=text,
        blocks=tuple(blocks or ()),
        source="libreoffice_pdf_text",
    )


def extract_office_born_digital(content: bytes, mime_type: str) -> OfficeExtractResult | None:
    global _BLOCK_INDEX  # noqa: PLW0603
    _BLOCK_INDEX = 0
    base = normalize_mime(mime_type)
    try:
        if base in LEGACY_OFFICE_MIMES:
            lo = _extract_via_libreoffice_to_pdf(content, base)
            if lo is not None:
                return lo
        if base == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return _extract_docx(content)
        if base == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            return _extract_xlsx(content)
        if base == "application/vnd.ms-excel":
            return _extract_via_libreoffice_to_pdf(content, base)
        if base == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            return _extract_pptx(content)
        if base == "application/vnd.ms-powerpoint":
            lo = _extract_via_libreoffice_to_pdf(content, base)
            return lo or _extract_pptx(content)
        if base == "application/msword":
            lo = _extract_via_libreoffice_to_pdf(content, base)
            if lo:
                return lo
        if base == "application/vnd.oasis.opendocument.text":
            return _extract_odt(content)
        if base == "application/vnd.oasis.opendocument.spreadsheet":
            return _extract_ods(content)
        if base == "application/vnd.oasis.opendocument.presentation":
            return _extract_odp(content)
        if base in {"application/rtf", "text/rtf"}:
            return _extract_rtf(content)
        if base == "text/csv":
            return _extract_csv(content)
        if base == "text/plain":
            return _extract_plain(content, source="plain")
        if base == "text/markdown":
            return _extract_plain(content, source="markdown")
        if base == "text/html":
            return _extract_html(content)
        if base == "message/rfc822":
            return _extract_eml(content)
        if base == "application/vnd.ms-outlook":
            return _extract_msg(content)
    except Exception as exc:
        logger.warning("office_extract_failed mime=%s err=%s", base, type(exc).__name__)
        return None
    return None


def build_heading_layout(blocks: tuple[ExtractionBlock, ...]) -> LayoutIrDocument | None:
    if not blocks:
        return None
    ir_blocks: list[LayoutIrBlock] = []
    for block in blocks:
        ir_blocks.append(
            LayoutIrBlock(
                page=block.page,
                x=block.x,
                y=block.y,
                width=block.width,
                height=block.height,
                text=block.text,
                block_index=block.block_index,
            )
        )
    page = LayoutIrPage(
        page=1,
        width_pt=612.0,
        height_pt=792.0,
        blocks=tuple(ir_blocks),
        lines=(),
        tables=(),
        vectors=(),
        widgets=(),
    )
    return LayoutIrDocument(version=LAYOUT_IR_VERSION, pages=(page,))
